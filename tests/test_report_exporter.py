"""
Tests for the institutional-design PDF/PPTX/Excel export builders
(modules/report_exporter.py) — Property Analysis tab review pass.

These are smoke tests: given representative fixtures (including a full
underwriting block, an IC summary, and a business plan), each builder
must return non-empty, correctly-signed bytes without raising, across a
range of "missing optional data" shapes (no ic_summary, no plan, no
underwriting, waterfall vs. simple equity structure).
"""

import io

import pytest

from modules.report_exporter import (
    build_pdf_report, build_pptx_report, build_excel_workbook, compose_fallback_ic_summary,
    build_owner_teaser_pdf,
)


def _prop(**overrides) -> dict:
    base = {
        "address": "123 Main St",
        "bbl": "1000010001",
        "borough": "Manhattan",
        "block": "1",
        "lot": "1",
        "lot_sf": 5000,
        "far_built": 2.0,
        "far_max": 6.0,
        "unused_far_pct": 66.7,
        "zoning_dist": "R7A",
        "strategies": ["Demolition"],
        "owner": "123 Main St LLC",
        "owner_type": "LLC",
        "last_sale_price": "$3,500,000",
        "last_sale_date": "2021-05-01",
        "distress_signal": "No Signal",
        "deal_score": {"score": 82, "tier": "Strong Lead"},
        "opportunity": {"score": 70},
        "market_comps": {"median_price_psf": 850, "count": 12},
        "pipeline": {"count": 3, "total_units": 240},
        "business_plan": {
            "acquisition": {"estimate": 3_500_000, "basis": "Comparable land sales"},
            "total_dev_cost": 12_000_000,
            "profit": 2_000_000,
            "margin_pct": 16.7,
            "bullets": ["Demolish existing 2-story building", "Build to max FAR"],
        },
        "underwriting": {
            "scenario_label": "Ground-Up Mixed-Use",
            "irr": 0.184,
            "equity_multiple": 2.1,
            "total_dev_cost": 12_000_000,
            "year1_noi": 780_000,
            "equity_structure": "simple",
        },
    }
    base.update(overrides)
    return base


def _ic_summary() -> dict:
    return {
        "recommendation": "GO",
        "thesis": ["Strong rezoning upside", "Below-market basis"],
        "key_risks": ["ULURP timeline risk", "Construction cost inflation"],
        "next_steps": ["Order title report", "Confirm zoning with DOB"],
    }


# ── Fallback IC summary ──────────────────────────────────────────────────────

def test_fallback_ic_summary_strong_lead_recommends_go():
    result = compose_fallback_ic_summary(_prop(deal_score={"score": 82, "tier": "Strong Lead"}))
    assert result["recommendation"] == "GO"
    assert result["thesis"]
    assert result["key_risks"]
    assert result["next_steps"]


def test_fallback_ic_summary_watch_tier_recommends_watch():
    result = compose_fallback_ic_summary(_prop(deal_score={"score": 55, "tier": "Watch"}))
    assert result["recommendation"] == "WATCH"


def test_fallback_ic_summary_pass_tier_recommends_reject():
    result = compose_fallback_ic_summary(_prop(deal_score={"score": 15, "tier": "Pass"}))
    assert result["recommendation"] == "REJECT"


def test_fallback_ic_summary_distress_signal_surfaces_as_risk():
    result = compose_fallback_ic_summary(_prop(distress_signal="Strong Signal"))
    assert any("Strong Signal" in r for r in result["key_risks"])


def test_fallback_ic_summary_no_underwriting_prompts_next_step():
    result = compose_fallback_ic_summary(_prop(underwriting=None))
    assert any("Underwriting" in s for s in result["next_steps"])


def test_fallback_ic_summary_never_raises_on_empty_prop():
    result = compose_fallback_ic_summary({})
    assert result["recommendation"] in ("GO", "WATCH", "REJECT")
    assert result["thesis"]
    assert result["key_risks"]
    assert result["next_steps"]


def test_fallback_ic_summary_consumable_by_pdf_and_pptx_builders():
    # The whole point: this dict must be a drop-in ic_summary for both
    # exporters, which previously always received None from app.py.
    ic = compose_fallback_ic_summary(_prop())
    pdf_bytes = build_pdf_report(_prop(), ic)
    pptx_bytes = build_pptx_report(_prop(), ic)
    assert pdf_bytes[:4] == b"%PDF"
    assert pptx_bytes[:2] == b"PK"


# ── PDF ──────────────────────────────────────────────────────────────────────

def test_pdf_report_full_fixture_returns_nonempty_bytes():
    pdf_bytes = build_pdf_report(_prop(), _ic_summary())
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 500
    assert pdf_bytes.startswith(b"%PDF")


def test_pdf_report_minimal_prop_no_ic_summary():
    minimal = {"address": "1 Test Ave", "bbl": "1", "borough": "Queens"}
    pdf_bytes = build_pdf_report(minimal, None)
    assert pdf_bytes.startswith(b"%PDF")


def test_pdf_report_waterfall_equity_structure():
    prop = _prop(underwriting={
        "scenario_label": "Ground-Up Mixed-Use", "equity_structure": "waterfall",
        "total_dev_cost": 12_000_000, "year1_noi": 780_000,
        "lp_irr": 0.15, "gp_irr": 0.28, "total_gp_promote": 500_000,
        "irr": None, "equity_multiple": None,
    })
    pdf_bytes = build_pdf_report(prop, _ic_summary())
    assert pdf_bytes.startswith(b"%PDF")


# ── PPTX ─────────────────────────────────────────────────────────────────────

def test_pptx_report_full_fixture_returns_nonempty_bytes():
    pptx_bytes = build_pptx_report(_prop(), _ic_summary())
    assert isinstance(pptx_bytes, bytes)
    assert len(pptx_bytes) > 500
    # .pptx is a zip container
    assert pptx_bytes[:2] == b"PK"


def test_pptx_report_minimal_prop_no_optional_sections():
    minimal = {"address": "1 Test Ave", "bbl": "1", "borough": "Queens"}
    pptx_bytes = build_pptx_report(minimal, None)
    assert pptx_bytes[:2] == b"PK"


def test_pptx_report_waterfall_equity_structure():
    prop = _prop(underwriting={
        "scenario_label": "Ground-Up Mixed-Use", "equity_structure": "waterfall",
        "total_dev_cost": 12_000_000, "year1_noi": 780_000,
        "lp_irr": 0.15, "gp_irr": 0.28, "total_gp_promote": 500_000,
        "irr": None, "equity_multiple": None,
    })
    pptx_bytes = build_pptx_report(prop, _ic_summary())
    assert pptx_bytes[:2] == b"PK"


def test_pptx_report_no_recommendation_still_builds():
    prop = _prop()
    pptx_bytes = build_pptx_report(prop, {"thesis": [], "key_risks": [], "next_steps": []})
    assert pptx_bytes[:2] == b"PK"


# ── Excel ────────────────────────────────────────────────────────────────────

def test_excel_workbook_returns_nonempty_bytes():
    xlsx_bytes = build_excel_workbook([_prop(), _prop(address="456 Elm St", deal_score={"score": 55, "tier": "Watch"})])
    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 500
    assert xlsx_bytes[:2] == b"PK"


def test_excel_workbook_empty_properties_list():
    xlsx_bytes = build_excel_workbook([])
    assert xlsx_bytes[:2] == b"PK"


def test_excel_workbook_all_three_tiers_render():
    props = [
        _prop(address="A", deal_score={"score": 90, "tier": "Strong Lead"}),
        _prop(address="B", deal_score={"score": 60, "tier": "Watch"}),
        _prop(address="C", deal_score={"score": 20, "tier": "Pass"}),
    ]
    xlsx_bytes = build_excel_workbook(props)
    assert xlsx_bytes[:2] == b"PK"


def test_excel_workbook_no_pro_forma_comps_or_budget_sheets_without_data():
    # The base _prop() fixture's underwriting/market_comps have only
    # summary numbers, no annual_cash_flows/cost_breakdown/comps list.
    xlsx_bytes = build_excel_workbook([_prop()])
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert not any("Pro Forma" in n for n in wb.sheetnames)
    assert not any("Comps" in n for n in wb.sheetnames)
    assert not any("Construction Budget" in n for n in wb.sheetnames)


def test_excel_workbook_pro_forma_and_budget_sheets_created_when_present():
    prop = _prop(underwriting={
        "scenario_label": "Ground-Up Mixed-Use",
        "irr": 0.184, "equity_multiple": 2.1,
        "total_dev_cost": 12_000_000, "year1_noi": 780_000,
        "equity_structure": "simple",
        "annual_cash_flows": [
            {"year": 0, "phase": "construction", "noi": 0, "debt_service": 0,
             "reversion_proceeds": 0, "equity_cf": -4_000_000, "unlevered_cf": -12_000_000},
            {"year": 1, "phase": "exit", "noi": 780_000, "debt_service": 400_000,
             "reversion_proceeds": 15_000_000, "equity_cf": 15_380_000, "unlevered_cf": 15_780_000},
        ],
        "cost_breakdown": {
            "acquisition_cost": 3_500_000, "closing_cost": 105_000,
            "hard_cost": 6_800_000, "soft_cost": 1_360_000,
            "contingency": 816_500, "total_dev_cost": 12_581_500,
        },
    })
    xlsx_bytes = build_excel_workbook([prop])
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert any("Pro Forma" in n for n in wb.sheetnames)
    assert any("Construction Budget" in n for n in wb.sheetnames)
    pf_sheet = next(wb[n] for n in wb.sheetnames if "Pro Forma" in n)
    assert pf_sheet.cell(row=4, column=1).value == 0  # year 0
    assert pf_sheet.cell(row=5, column=1).value == 1  # year 1
    budget_sheet = next(wb[n] for n in wb.sheetnames if "Construction Budget" in n)
    assert budget_sheet.cell(row=3, column=2).value == 3_500_000  # acquisition cost


def test_excel_workbook_construction_budget_sheet_without_trade_budget_unchanged():
    # No "trade_budget" key present (every property before this batch, and
    # every property today unless the new expander was opened) — the
    # existing 5-line summary must render exactly as before, nothing extra.
    prop = _prop(underwriting={
        "scenario_label": "Ground-Up Mixed-Use", "irr": 0.184, "equity_multiple": 2.1,
        "total_dev_cost": 12_000_000, "year1_noi": 780_000, "equity_structure": "simple",
        "cost_breakdown": {
            "acquisition_cost": 3_500_000, "closing_cost": 105_000,
            "hard_cost": 6_800_000, "soft_cost": 1_360_000,
            "contingency": 816_500, "total_dev_cost": 12_581_500,
        },
    })
    xlsx_bytes = build_excel_workbook([prop])
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    budget_sheet = next(wb[n] for n in wb.sheetnames if "Construction Budget" in n)
    assert budget_sheet.cell(row=3, column=1).value == "Acquisition Cost"
    assert budget_sheet.cell(row=8, column=1).value == "Total Development Cost"
    assert budget_sheet.cell(row=9, column=1).value is None  # nothing appended


def test_excel_workbook_construction_budget_sheet_appends_trade_detail_when_present():
    from modules.construction_budget_estimator import estimate_trade_level_budget
    trade_budget = estimate_trade_level_budget(10_000, budget_tier="Standard")
    prop = _prop(underwriting={
        "scenario_label": "Ground-Up Mixed-Use", "irr": 0.184, "equity_multiple": 2.1,
        "total_dev_cost": 12_000_000, "year1_noi": 780_000, "equity_structure": "simple",
        "cost_breakdown": {
            "acquisition_cost": 3_500_000, "closing_cost": 105_000,
            "hard_cost": 6_800_000, "soft_cost": 1_360_000,
            "contingency": 816_500, "total_dev_cost": 12_581_500,
        },
        "trade_budget": trade_budget,
    })
    xlsx_bytes = build_excel_workbook([prop])
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    budget_sheet = next(wb[n] for n in wb.sheetnames if "Construction Budget" in n)
    # Existing 5-line summary still present, unchanged, at the same rows.
    assert budget_sheet.cell(row=3, column=1).value == "Acquisition Cost"
    assert budget_sheet.cell(row=8, column=1).value == "Total Development Cost"
    # Trade detail appended below it.
    sheet_text = " ".join(
        str(budget_sheet.cell(row=r, column=1).value or "")
        for r in range(1, budget_sheet.max_row + 1)
    )
    assert "Trade-Level Detail" in sheet_text
    assert "Structure & Superstructure" in sheet_text


def test_excel_workbook_comps_sheet_created_when_present():
    prop = _prop(market_comps={
        "median_price_psf": 850, "count": 1,
        "comps": [{"address": "1 Test Ave", "price": 3_000_000, "sqft": 5000,
                    "price_psf": 600.0, "asset_type": "Residential Sale", "date": "2024-01-01"}],
    })
    xlsx_bytes = build_excel_workbook([prop])
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert any("Comps" in n for n in wb.sheetnames)
    comps_sheet = next(wb[n] for n in wb.sheetnames if "Comps" in n)
    assert comps_sheet.cell(row=4, column=1).value == "1 Test Ave"


# ── included_sections: backward compatibility + new sections ────────────────
#
# Content-based comparisons (not raw byte equality) are used throughout —
# reportlab/openpyxl embed a creation timestamp that can differ by a second
# between two calls in the same test, which would make raw byte-equality
# assertions flaky without actually indicating a real content difference.

def _pdf_text(pdf_bytes: bytes) -> str:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)


def _pptx_text(pptx_bytes: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(pptx_bytes))
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


def _xlsx_dump(xlsx_bytes: bytes) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    return {name: [[c.value for c in row] for row in wb[name].iter_rows()] for name in wb.sheetnames}


def _composite_distress() -> dict:
    return {
        "score": 62, "tier": "Elevated",
        "breakdown": {
            "acris": {"score": 20, "max": 25, "weight": 0.25,
                      "reasoning": "1 foreclosure/lis pendens filing(s), 2 open lien(s)"},
            "hpd": {"score": 10, "max": 20, "weight": 0.20, "reasoning": "3 open HPD violation(s)"},
        },
        "components_with_data": ["acris", "hpd"],
    }


def test_pdf_report_included_sections_none_matches_default_behavior():
    prop = _prop()
    ic = _ic_summary()
    baseline = _pdf_text(build_pdf_report(prop, ic))
    explicit_none = _pdf_text(build_pdf_report(prop, ic, included_sections=None))
    assert baseline == explicit_none


def test_pptx_report_included_sections_none_matches_default_behavior():
    prop = _prop()
    ic = _ic_summary()
    baseline = _pptx_text(build_pptx_report(prop, ic))
    explicit_none = _pptx_text(build_pptx_report(prop, ic, included_sections=None))
    assert baseline == explicit_none


def test_excel_workbook_included_sections_none_matches_default_behavior():
    props = [_prop()]
    baseline = _xlsx_dump(build_excel_workbook(props))
    explicit_none = _xlsx_dump(build_excel_workbook(props, included_sections=None))
    assert baseline == explicit_none


def test_pdf_report_included_sections_excludes_thesis_when_false():
    prop = _prop()
    ic = _ic_summary()
    included = build_pdf_report(prop, ic, included_sections={"thesis": True})
    excluded = build_pdf_report(prop, ic, included_sections={"thesis": False})
    assert "Investment Thesis" in _pdf_text(included)
    assert "Investment Thesis" not in _pdf_text(excluded)


def test_pdf_report_included_sections_excludes_ownership_when_false():
    prop = _prop()
    included = build_pdf_report(prop, None, included_sections={"ownership": True})
    excluded = build_pdf_report(prop, None, included_sections={"ownership": False})
    assert "123 Main St LLC" in _pdf_text(included)
    assert "123 Main St LLC" not in _pdf_text(excluded)


def test_pptx_report_included_sections_excludes_business_plan_when_false():
    prop = _prop()
    included = build_pptx_report(prop, None, included_sections={"business_plan": True})
    excluded = build_pptx_report(prop, None, included_sections={"business_plan": False})
    assert "Preliminary Acquisition Estimate" in _pptx_text(included)
    assert "Preliminary Acquisition Estimate" not in _pptx_text(excluded)


def test_pdf_report_new_distress_section_renders_when_present():
    prop = _prop(composite_distress=_composite_distress())
    text = _pdf_text(build_pdf_report(prop, None, included_sections={"distress": True}))
    assert "Distress Signals" in text
    assert "62/100" in text
    assert "ACRIS" in text.upper()


def test_pdf_report_new_tax_abatement_rent_stab_section_renders_when_present():
    prop = _prop(
        tax_abatement={"program": "421-a", "summary": "Partial exemption"},
        rent_stab_signal={"likely_stabilized": True, "confidence": 0.97},
    )
    text = _pdf_text(build_pdf_report(prop, None, included_sections={"tax_abatement_rent_stab": True}))
    assert "Tax Abatement & Rent Stabilization" in text
    assert "421-a" in text
    assert "Likely Rent Stabilized" in text


def test_pdf_report_new_market_comps_section_renders_when_present():
    prop = _prop(market_comps={
        "median_price_psf": 850, "count": 1,
        "comps": [{"address": "1 Test Ave", "price": 3_000_000, "sqft": 5000, "price_psf": 600.0}],
    })
    text = _pdf_text(build_pdf_report(prop, None, included_sections={"market_comps": True}))
    assert "Market Comps" in text
    assert "1 Test Ave" in text


def test_pptx_report_new_sections_render_when_included_sections_provided():
    prop = _prop(
        composite_distress=_composite_distress(),
        tax_abatement={"program": "421-a", "summary": "Partial exemption"},
        rent_stab_signal={"likely_stabilized": True, "confidence": 0.97},
    )
    text = _pptx_text(build_pptx_report(prop, None, included_sections={
        "distress": True, "tax_abatement_rent_stab": True, "market_comps": True,
    }))
    assert "Distress Signals" in text
    assert "Tax Abatement & Rent Stabilization" in text


def test_excel_workbook_included_sections_excludes_ownership_when_false():
    prop = _prop()
    included = _xlsx_dump(build_excel_workbook([prop], included_sections={"ownership": True}))
    excluded = _xlsx_dump(build_excel_workbook([prop], included_sections={"ownership": False}))
    detail_sheet_name = next(n for n in included if n not in ("Summary",))
    included_text = " ".join(str(c) for row in included[detail_sheet_name] for c in row)
    excluded_text = " ".join(str(c) for row in excluded[detail_sheet_name] for c in row)
    assert "123 Main St LLC" in included_text
    assert "123 Main St LLC" not in excluded_text


# ── Backward-compatibility regression: the Site Finder call site ────────────
#
# modules/site_finder_ui.py calls build_pdf_report(prop, ic) / build_pptx_report(prop, ic)
# with exactly two positional args, never included_sections. These tests lock
# in that exact calling convention and prove the new sections never appear
# there — regardless of what extra keys (composite_distress, market_comps)
# happen to be present on its `prop` dict — since Site Finder's own
# `distress` key (a *different* shape, from modules/ownership_research.py)
# and `market_comps` key are already populated on every property it exports.

def test_pdf_report_backward_compatible_with_site_finder_call_signature():
    prop = _prop(
        composite_distress=_composite_distress(),
        tax_abatement={"program": "421-a", "summary": "Partial exemption"},
        # Site Finder's own differently-shaped "distress" key (never read by
        # the new composite-distress section, which reads "composite_distress").
        distress={"level": "Moderate", "score": 6, "evidence": ["1 open lien"]},
    )
    ic = _ic_summary()
    pdf_bytes = build_pdf_report(prop, ic)  # exact Site Finder call signature
    assert pdf_bytes.startswith(b"%PDF")
    text = _pdf_text(pdf_bytes)
    assert "Distress Signals" not in text
    assert "Tax Abatement & Rent Stabilization" not in text


def test_pptx_report_backward_compatible_with_site_finder_call_signature():
    prop = _prop(
        composite_distress=_composite_distress(),
        tax_abatement={"program": "421-a", "summary": "Partial exemption"},
        distress={"level": "Moderate", "score": 6, "evidence": ["1 open lien"]},
    )
    ic = _ic_summary()
    pptx_bytes = build_pptx_report(prop, ic)  # exact Site Finder call signature
    assert pptx_bytes[:2] == b"PK"
    text = _pptx_text(pptx_bytes)
    assert "Distress Signals" not in text
    assert "Tax Abatement & Rent Stabilization" not in text


def test_excel_workbook_backward_compatible_with_site_finder_call_signature():
    prop = _prop(
        composite_distress=_composite_distress(),
        distress={"level": "Moderate", "score": 6, "evidence": ["1 open lien"]},
    )
    xlsx_bytes = build_excel_workbook([prop])  # exact Site Finder call signature
    assert xlsx_bytes[:2] == b"PK"


# ── Owner Outreach Teaser PDF (Batch E, Feature 7) ───────────────────────────
#
# Owner-facing document — must NEVER leak internal deal economics (Deal
# Score, distress signal, margin, IRR), unlike build_pdf_report() which is
# explicitly framed "CONFIDENTIAL — INVESTMENT SCREENING MEMORANDUM".

def test_owner_teaser_pdf_returns_pdf_magic_bytes():
    pdf_bytes = build_owner_teaser_pdf(_prop())
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")


def test_owner_teaser_pdf_contains_owner_safe_fields():
    text = _pdf_text(build_owner_teaser_pdf(_prop()))
    assert "123 Main St" in text
    assert "67%" in text or "66.7%" in text or "67" in text  # unused FAR
    assert "3,500,000" in text  # acquisition estimate from business_plan.acquisition.estimate


def test_owner_teaser_pdf_never_leaks_deal_economics():
    prop = _prop(composite_distress=_composite_distress())
    text = _pdf_text(build_owner_teaser_pdf(prop))
    assert "Deal Score" not in text
    assert "82" not in text          # deal_score.score
    assert "Strong Lead" not in text  # deal_score.tier
    assert "Distress" not in text
    assert "Margin" not in text
    assert "16.7%" not in text        # business_plan.margin_pct
    assert "IRR" not in text


def test_owner_teaser_pdf_handles_missing_fields_gracefully():
    # Never raises on missing/malformed prop fields — only on the
    # reportlab-ImportError path.
    pdf_bytes = build_owner_teaser_pdf({})
    assert pdf_bytes.startswith(b"%PDF")
    text = _pdf_text(pdf_bytes)
    assert "—" in text  # graceful "—" fallback for unavailable fields


def test_owner_teaser_pdf_handles_malformed_business_plan():
    prop = _prop(business_plan="not a dict", unused_far_pct=None)
    pdf_bytes = build_owner_teaser_pdf(prop)
    assert pdf_bytes.startswith(b"%PDF")


def test_owner_teaser_pdf_raises_import_error_when_reportlab_missing(monkeypatch):
    import sys
    monkeypatch.setitem(sys.modules, "reportlab", None)
    monkeypatch.setitem(sys.modules, "reportlab.lib.pagesizes", None)
    with pytest.raises(ImportError):
        build_owner_teaser_pdf(_prop())
